# -*- coding: utf-8 -*-
"""세로형(A4) 맞춤 제안서 PPTX 빌더 — python-pptx.

모든 슬라이드: 상단 우측 = 리드 고객사명, 하단 우측 = 클라우다이크 로고, 하단 좌측 = 쪽번호.
본문 블록은 PowerPoint 도형(둥근 사각형 카드·콜아웃·스탯 타일·단계 원형 번호)으로 그린다.
위에서 아래로 쌓는 커서 방식이며, 슬라이드 하단 여백을 넘기면 경고를 출력한다 → 내용을 줄이거나 슬라이드를 나눈다.

    import os, sys; sys.path.insert(0, "<이 스킬>/scripts")
    from pptx_builder import Deck
    deck = Deck("(주)에이아이랩스")
    deck.cover("외부 협업 파일 보안 공유 체계 구축 제안서", "기업 드라이브 전환 — 링크 정책·접근 이력·10TB 이관",
               "2026년 10월 29일", "CLOUDIKE Inc.", contact="영업팀 하정우")
    deck.toc()                                   # 2쪽: 제안 목차 (장 제목 + 쪽번호 자동)
    s = deck.slide("제안 배경 및 필요성", kicker="01 · BACKGROUND", section="제안 배경 및 필요성")
    s.callout("핵심 제안 명제 …")
    s.stats([("150계정 · 10TB", "도입 규모"), ("12월 초", "계약"), ("[견적 협의]", "첫해 금액")])
    s.cards([("외부 링크 정책", "만료일·비밀번호 필수", "기본 제공")], cols=2)
    s.table([["#", "고객 원문", "대응", "제공"], ["R1", "…", "…", "기본 제공"]], widths=[0.5, 3, 3, 1], badge_col=3)
    s.steps([("체험", "11/9–11/23", "연구팀 일부, 링크 정책 시험")])
    deck.save("제안서.pptx")

텍스트 안의 **굵게** 지원. 색은 클라우다이크 로고 앰버(FAB329) + 다크 네이비.
"""
import math, os, re, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, "..", "assets", "cloudike_logo.png")

# 팔레트 — 앰버는 강조 1곳에만, 네이비가 지배색
NAVY = "1F2A44"; AMBER = "FAB329"; AMBER_DK = "B7791F"; TINT = "FFF6E0"
INK = "1F2937"; GRAY = "6B7280"; LINE = "E5E7EB"; SOFT = "F5F7FA"; WHITE = "FFFFFF"
BADGES = {"기본 제공": ("E6F4EF", "0F766E"), "시연 확인": ("FFF1D6", "B45309"),
          "확인 필요": ("FDECEC", "B42318"), "가정": ("EEF0F5", "475467"), "협의": ("EEF0F5", "475467")}
# 한 칸에 여러 등급이 섞이면(예: '기본 제공 + 시연 확인') 더 보수적인 등급의 색을 쓴다 — B 를 A 처럼 보이게 하지 않기 위해
BADGE_PRIORITY = ["확인 필요", "시연 확인", "가정", "협의", "기본 제공"]


def badge_style(label):
    for k in BADGE_PRIORITY:
        if k in str(label):
            return BADGES[k]
    return None

W, H = 7.5, 10.833            # PowerPoint 'A4 용지' 세로
MX = 0.55                     # 좌우 여백
TOP = 0.95                    # 본문 시작(제목 위치)
BOTTOM = H - 0.85             # 본문 하한 (로고·쪽번호 영역 위)
CW = W - 2 * MX               # 본문 폭


def _rgb(h): return RGBColor.from_string(h)


def _ea(run, font):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)


def _em(ch):
    o = ord(ch)
    if o >= 0x1100 and not (0x2000 <= o <= 0x206F): return 1.0   # 한글·CJK·기호
    if ch == " ": return 0.32
    if ch in "il.,:;'|!": return 0.3
    if ch.isupper() or ch in "mwMW@%": return 0.68
    return 0.56


# ---------- 실제 폰트 폭 측정
# PowerPoint 는 한글을 글자 단위로 줄바꿈한다(「단/위」). 그래서 빌더가 폰트 파일로 글자 폭을 재서
# 단어 경계에서 직접 줄을 나누고(<a:br/>), 같은 계산으로 상자 높이를 정한다.
# 폰트 파일을 못 찾으면 어림값으로 높이만 추정하고 줄바꿈은 PowerPoint 에 맡긴다.
FONT = {"name": "Pretendard"}          # Deck 이 설정
FIT = 0.95                             # 우리가 나눈 줄을 PowerPoint 가 다시 쪼개지 않게 폭 여유
_FONTS, _CWS = {}, {}
_ALIAS = {"맑은 고딕": "malgun", "malgun gothic": "malgun", "나눔고딕": "nanumgothic"}


def _font_file(name, bold):
    key = _ALIAS.get(name.lower(), _ALIAS.get(name, name)).replace(" ", "").lower()
    dirs = [os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts"),
            os.path.join(os.environ.get("LOCALAPPDATA", ""), "Microsoft", "Windows", "Fonts"),
            "/Library/Fonts", os.path.expanduser("~/Library/Fonts"), "/usr/share/fonts", os.path.expanduser("~/.fonts")]
    files = []
    for d in dirs:
        if d and os.path.isdir(d):
            for root, _, fs in os.walk(d):
                files += [os.path.join(root, f) for f in fs
                          if f.lower().endswith((".ttf", ".otf")) and key in f.lower().replace("-", "").replace(" ", "")]
    if not files:
        return None

    def score(f):
        n = os.path.basename(f).lower()
        is_bold = (re.search(r"bold", n) and not re.search(r"extrabold|semibold", n)) or n.endswith("bd.ttf")
        if bold:
            return 0 if is_bold else (1 if "semibold" in n else 2)
        return 0 if "regular" in n else (1 if not re.search(r"bold|light|thin|black|medium|bd\.ttf", n) else 2)
    return sorted(files, key=score)[0]


def _font(bold):
    k = (FONT["name"], bold)
    if k not in _FONTS:
        try:
            from PIL import ImageFont
            path = _font_file(FONT["name"], bold)
            _FONTS[k] = ImageFont.truetype(path, 100) if path else None
        except Exception:
            _FONTS[k] = None
    return _FONTS[k]


def _cw(ch, bold):
    """글자 폭 (em)."""
    k = (ch, bold, FONT["name"])
    if k not in _CWS:
        f = _font(bold)
        _CWS[k] = f.getlength(ch) / 100 if f else _em(ch) * (1.05 if bold else 1.0)
    return _CWS[k]


def wrap(text, width_in, size, bold=False):
    """단어 경계 줄바꿈. 반환: 문단마다 [줄마다 [(글자, 굵기), ...]]. '**굵게**' 표기 처리."""
    limit = max(0.5, width_in * 72 / size * FIT)
    paras = []
    for para in str(text).split("\n"):
        chars = []
        for j, seg in enumerate(para.split("**")):
            chars += [(c, bold or j % 2 == 1) for c in seg]
        toks, cur = [], []
        for c in chars:
            if c[0] == " ":
                if cur: toks.append(cur); cur = []
                toks.append([c])
            else:
                cur.append(c)
        if cur: toks.append(cur)
        lines, w = [[]], 0.0
        for t in toks:
            tw = sum(_cw(*c) for c in t)
            if t[0][0] == " ":
                if lines[-1]: lines[-1].append(t[0]); w += tw
                continue
            if w + tw > limit and lines[-1]:
                while lines[-1] and lines[-1][-1][0] == " ": lines[-1].pop()
                lines.append([]); w = 0.0
            if tw > limit:                       # 한 단어가 한 줄보다 길면 글자 단위
                for c in t:
                    if w + _cw(*c) > limit and lines[-1]: lines.append([]); w = 0.0
                    lines[-1].append(c); w += _cw(*c)
            else:
                lines[-1].extend(t); w += tw
        while lines[-1] and lines[-1][-1][0] == " ": lines[-1].pop()
        paras.append(lines)
    return paras


def est_lines(text, width_in, size, bold=False):
    """폭 width_in(인치)에서 size pt 텍스트의 줄 수 — 폰트가 있으면 실측, 없으면 보수적 어림."""
    if _font(False) is not None:
        return sum(len(lines) for lines in wrap(text, width_in, size, bold))
    cap = max(1.0, width_in * 72 / size)
    n = 0
    for para in str(text).replace("**", "").split("\n"):
        n += max(1, math.ceil(sum(_em(c) for c in para) * 1.12 / cap))
    return n


def est_h(text, width_in, size, spacing=1.32, bold=False):
    return est_lines(text, width_in, size, bold) * size * spacing / 72


def _shadow(shape, alpha=14):
    spPr = shape._element.spPr
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad="76200", dist="19050", dir="5400000",
                          algn="t", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
    etree.SubElement(clr, qn("a:alpha"), val=str(alpha * 1000))


class Slide:
    def __init__(self, deck, s, no):
        self.deck, self.s, self.no, self.y = deck, s, no, TOP
        self.font = deck.font

    # ---------- 기본 도구
    def _tb(self, x, y, w, h, text, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP, spacing=1.32, margin=0.0, name=None):
        tb = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name: tb.name = name
        self._fill(tb.text_frame, text, size, color, bold, align, anchor, spacing, margin, width=w - 2 * margin)
        return tb

    def _run(self, p, txt, size, bold, color):
        r = p.add_run(); r.text = txt
        r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = _rgb(color); _ea(r, self.font)

    def _fill(self, tf, text, size, color, bold, align, anchor, spacing=1.32, margin=0.0, width=None):
        tf.word_wrap = True; tf.vertical_anchor = anchor
        for side in ("left", "right", "top", "bottom"):
            setattr(tf, f"margin_{side}", Inches(margin))
        if width and _font(False) is not None:      # 단어 단위 줄바꿈을 직접 넣는다
            for i, lines in enumerate(wrap(text, width, size, bold)):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = align; p.line_spacing = spacing
                for li, line in enumerate(lines):
                    if li: p.add_line_break()
                    k = 0
                    while k < len(line):
                        b = line[k][1]; j = k
                        while j < len(line) and line[j][1] == b: j += 1
                        self._run(p, "".join(c for c, _ in line[k:j]), size, b, color); k = j
            return
        for i, para in enumerate(str(text).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align; p.line_spacing = spacing
            for j, seg in enumerate(para.split("**")):
                if not seg: continue
                r = p.add_run(); r.text = seg
                r.font.size = Pt(size); r.font.bold = bold or (j % 2 == 1)
                r.font.color.rgb = _rgb(color); _ea(r, self.font)

    def _box(self, x, y, w, h, fill=WHITE, line=None, radius=0.06, shadow=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=None):
        b = self.s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if name: b.name = name
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            b.adjustments[0] = radius
        b.fill.solid(); b.fill.fore_color.rgb = _rgb(fill)
        if line: b.line.color.rgb = _rgb(line); b.line.width = Pt(0.75)
        else: b.line.fill.background()
        if shadow: _shadow(b)
        else: b.shadow.inherit = False   # 기본 테마의 도형 그림자 끄기
        b.text_frame.text = ""
        return b

    def _advance(self, h, gap=0.22):
        self.y += h + gap
        if self.y - gap > BOTTOM + 0.02:
            self.deck.over[self.no] = self.y - gap - BOTTOM

    def remaining(self):
        return BOTTOM - self.y

    def badge(self, x, y, label, size=9):
        bg, fg = badge_style(label) or ("EEF0F5", "475467")
        w = sum(_em(c) for c in label) * size / 72 + 0.22
        b = self._box(x, y, w, size / 72 + 0.12, fill=bg, radius=0.5)
        self._fill(b.text_frame, label, size, fg, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 1.0, 0.02)
        return w

    # ---------- 본문 블록
    def text(self, t, size=12, color=INK, bold=False, gap=0.16):
        h = est_h(t, CW, size)
        self._tb(MX, self.y, CW, h, t, size, color, bold)
        self._advance(h, gap)

    def subhead(self, t, size=14):
        """슬라이드 안의 소제목 (앰버 점 + 굵은 텍스트)."""
        d = self.s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(MX), Inches(self.y + 0.07), Inches(0.12), Inches(0.12))
        d.fill.solid(); d.fill.fore_color.rgb = _rgb(AMBER); d.line.fill.background(); d.shadow.inherit = False
        h = est_h(t, CW - 0.22, size, bold=True)
        self._tb(MX + 0.22, self.y, CW - 0.22, h, t, size, NAVY, True)
        self._advance(h, 0.12)

    def callout(self, t, size=13, fill=TINT, color=INK):
        """핵심 메시지 박스 — 앰버 틴트 둥근 사각형, 굵은 따옴표 아이콘."""
        pad = 0.26
        h = est_h(t, CW - 2 * pad - 0.3, size, 1.4) + 2 * pad
        self._box(MX, self.y, CW, h, fill=fill, radius=0.08, name="callout")
        q = self._tb(MX + 0.16, self.y + 0.08, 0.4, 0.5, "“", 30, AMBER, True)
        self._tb(MX + pad + 0.3, self.y + pad, CW - 2 * pad - 0.3, h - 2 * pad, t, size, color, False, spacing=1.4)
        self._advance(h)

    def stats(self, items, size=19):
        """큰 숫자 타일 (값, 라벨[, 보조]) 2~4개 가로 배치."""
        n = len(items); gap = 0.18; w = (CW - gap * (n - 1)) / n
        h = max(est_h(it[0], w - 0.3, size, 1.15, bold=True) for it in items) + 0.5 + \
            max((est_h(it[2], w - 0.3, 9.5) if len(it) > 2 else 0) for it in items)
        for i, it in enumerate(items):
            x = MX + i * (w + gap)
            self._box(x, self.y, w, h, fill=NAVY if i == 0 else SOFT, radius=0.1, name=f"stat{i+1}")
            fg = WHITE if i == 0 else NAVY
            self._tb(x + 0.15, self.y + 0.14, w - 0.3, 0.22, it[1], 9.5, AMBER if i == 0 else GRAY, True)
            vh = est_h(it[0], w - 0.3, size, 1.15, bold=True)
            self._tb(x + 0.15, self.y + 0.38, w - 0.3, vh, it[0], size, fg, True, spacing=1.15)
            if len(it) > 2:
                self._tb(x + 0.15, self.y + 0.42 + vh, w - 0.3, h - vh - 0.5, it[2], 9.5, "D1D5DB" if i == 0 else GRAY)
        self._advance(h)

    def cards(self, items, cols=2, size=10.5, title_size=12):
        """카드 그리드: (제목, 본문[, 배지]) — 흰 카드 + 옅은 그림자."""
        gap = 0.18; w = (CW - gap * (cols - 1)) / cols; pad = 0.18
        rows = [items[i:i + cols] for i in range(0, len(items), cols)]
        for row in rows:
            hs = []
            for it in row:
                hh = est_h(it[0], w - 2 * pad, title_size, 1.2, bold=True) + 0.08 + est_h(it[1], w - 2 * pad, size, 1.3)
                if len(it) > 2 and it[2]: hh += 0.3
                hs.append(hh + 2 * pad)
            h = max(hs)
            for i, it in enumerate(row):
                x = MX + i * (w + gap); y = self.y
                self._box(x, y, w, h, fill=WHITE, line=LINE, radius=0.07, shadow=True, name="card")
                yy = y + pad
                if len(it) > 2 and it[2]:
                    self.badge(x + pad, yy, it[2]); yy += 0.3
                th = est_h(it[0], w - 2 * pad, title_size, 1.2, bold=True)
                self._tb(x + pad, yy, w - 2 * pad, th, it[0], title_size, NAVY, True, spacing=1.2)
                self._tb(x + pad, yy + th + 0.08, w - 2 * pad, h - (yy - y) - th - pad, it[1], size, INK, spacing=1.3)
            self._advance(h, gap if row is not rows[-1] else 0.22)

    def rows(self, items, size=10.5, label_w=1.55):
        """아이콘형 행 목록: (라벨, 본문[, 배지]) — 왼쪽 네이비 칩, 오른쪽 설명."""
        for it in items:
            bw = CW - label_w - 0.15
            h = max(est_h(it[1], bw - 0.3 - (0.0), size), 0.3) + 0.24 + (0.28 if len(it) > 2 and it[2] else 0)
            self._box(MX, self.y, CW, h, fill=SOFT, radius=0.08, name="row")
            chip = self._box(MX + 0.1, self.y + 0.1, label_w - 0.1, h - 0.2, fill=NAVY, radius=0.12)
            self._fill(chip.text_frame, it[0], 10.5, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 1.15, 0.05)
            yy = self.y + 0.12
            if len(it) > 2 and it[2]:
                self.badge(MX + label_w + 0.15, yy, it[2]); yy += 0.28
            self._tb(MX + label_w + 0.15, yy, bw - 0.15, h - (yy - self.y) - 0.1, it[1], size, INK, anchor=MSO_ANCHOR.TOP)
            self._advance(h, 0.1)
        self.y += 0.12

    def table(self, rows, widths=None, size=9.5, badge_col=None, header=True, center_cols=()):
        """표: 네이비 머리행 + 줄무늬. badge_col 열의 '기본 제공/시연 확인/확인 필요/가정'은 색 배지 텍스트."""
        nc = max(len(r) for r in rows); rows = [list(r) + [""] * (nc - len(r)) for r in rows]
        widths = widths or [1] * nc; tot = sum(widths); ws = [CW * w / tot for w in widths]
        pad = 0.07
        hs = [max(est_h(c, ws[j] - 2 * pad, size, 1.22, bold=header and i == 0) for j, c in enumerate(r)) + 2 * pad + 0.02
              for i, r in enumerate(rows)]
        h = sum(hs)
        gt = self.s.shapes.add_table(len(rows), nc, Inches(MX), Inches(self.y), Inches(CW), Inches(h))
        gt.name = "table"
        t = gt.table
        tblPr = t._tbl.tblPr
        sid = tblPr.find(qn("a:tableStyleId"))
        if sid is None: sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
        sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"   # No Style, Table Grid → 색은 직접 지정
        for j, wv in enumerate(ws): t.columns[j].width = Inches(wv)
        for i, r in enumerate(rows):
            t.rows[i].height = Inches(hs[i])
            for j, val in enumerate(r):
                c = t.cell(i, j)
                for side in ("left", "right", "top", "bottom"):
                    setattr(c, f"margin_{side}", Inches(pad))
                hdr = header and i == 0
                c.fill.solid(); c.fill.fore_color.rgb = _rgb(NAVY if hdr else (SOFT if i % 2 == 0 else WHITE))
                fg, bold = (WHITE, True) if hdr else (INK, False)
                if not hdr and badge_col is not None and j == badge_col and badge_style(val):
                    fg, bold = badge_style(val)[1], True
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = c.text_frame; tf.text = ""
                self._fill(tf, val, size, fg, bold,
                           PP_ALIGN.CENTER if (hdr or j in center_cols) else PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE, 1.22,
                           width=ws[j] - 2 * pad)
                # 셀 테두리: 옅은 회색 가로선
                tcPr = c._tc.get_or_add_tcPr()
                for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                    ln = etree.SubElement(tcPr, qn(tag), w="6350")
                    if tag in ("a:lnL", "a:lnR"):
                        etree.SubElement(ln, qn("a:noFill"))
                    else:
                        sf = etree.SubElement(ln, qn("a:solidFill")); etree.SubElement(sf, qn("a:srgbClr"), val=LINE)
                # tcPr 자식 순서: ln* 다음 fill 이어야 스키마에 맞음 → fill 요소를 맨 뒤로
                for f in tcPr.findall(qn("a:solidFill")):
                    tcPr.remove(f); tcPr.append(f)
        self._advance(h, 0.32)   # PowerPoint 가 행을 추정보다 키우는 경우가 있어 여유를 둔다

    def steps(self, items, size=10.5):
        """세로 단계 흐름: 번호 원 + 연결선 + (제목, 시점, 설명[, 통과 기준])."""
        x0 = MX + 0.22; bw = CW - 0.75
        tops = []
        for i, it in enumerate(items):
            body = it[2] + (f"\n**통과 기준** {it[3]}" if len(it) > 3 and it[3] else "")
            h = max(0.62, 0.42 + est_h(body, bw - 0.36, size, 1.3) + 0.16)
            tops.append((self.y, h))
            self._box(MX + 0.6, self.y, bw, h, fill=WHITE, line=LINE, radius=0.08, shadow=True, name=f"step{i+1}")
            self._tb(MX + 0.78, self.y + 0.12, bw - 1.9, 0.26, it[0], 12, NAVY, True)
            self._tb(MX + 0.6 + bw - 1.75, self.y + 0.13, 1.6, 0.24, it[1], 9.5, AMBER_DK, True, align=PP_ALIGN.RIGHT)
            self._tb(MX + 0.78, self.y + 0.42, bw - 0.36, h - 0.5, body, size, INK, spacing=1.3)
            self.y += h + 0.14
        # 연결선은 원보다 먼저 그려야 뒤로 깔림 → 원을 마지막에
        if len(tops) > 1:
            ln = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0 + 0.17), Inches(tops[0][0] + 0.2),
                                             Inches(x0 + 0.17), Inches(tops[-1][0] + 0.2))
            ln.line.color.rgb = _rgb(LINE); ln.line.width = Pt(2)
        for i, (ty, h) in enumerate(tops):
            c = self.s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x0), Inches(ty + 0.03), Inches(0.36), Inches(0.36))
            c.fill.solid(); c.fill.fore_color.rgb = _rgb(AMBER if i == 0 else NAVY); c.line.fill.background(); c.shadow.inherit = False
            self._fill(c.text_frame, str(i + 1), 11, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 1.0, 0.0)
        self.y += 0.08
        if self.y > BOTTOM + 0.02:
            self.deck.over[self.no] = self.y - BOTTOM

    def compare(self, left, right, size=10.5):
        """좌우 비교 카드: (머리, [줄...]) 두 개. 오른쪽(제안)이 강조."""
        gap = 0.2; w = (CW - gap) / 2; pad = 0.2
        hl = 0.5 + sum(est_h("· " + t, w - 2 * pad, size, 1.3) + 0.06 for t in left[1])
        hr = 0.5 + sum(est_h("· " + t, w - 2 * pad, size, 1.3) + 0.06 for t in right[1])
        h = max(hl, hr) + pad
        for k, (head, lines) in enumerate((left, right)):
            x = MX + k * (w + gap)
            self._box(x, self.y, w, h, fill=SOFT if k == 0 else TINT, radius=0.07, name="compare")
            self._tb(x + pad, self.y + 0.16, w - 2 * pad, 0.3, head, 12.5, GRAY if k == 0 else NAVY, True)
            yy = self.y + 0.52
            for t in lines:
                lh = est_h("· " + t, w - 2 * pad, size, 1.3)
                self._tb(x + pad, yy, w - 2 * pad, lh, "· " + t, size, INK, spacing=1.3); yy += lh + 0.06
        self._advance(h)

    def note(self, t, size=9):
        h = est_h("※ " + t, CW, size)
        self._tb(MX, self.y, CW, h, "※ " + t, size, GRAY)
        self._advance(h, 0.08)

    def gap(self, h=0.15):
        self.y += h

    def notes(self, t):
        self.s.notes_slide.notes_text_frame.text = t


class Deck:
    def __init__(self, customer, font="Pretendard", logo=LOGO, footer=None):
        self.p = Presentation()
        self.p.slide_width = Inches(W); self.p.slide_height = Inches(H)
        self.customer, self.font, self.logo = customer, font, logo
        FONT["name"] = font
        self.footer = footer or f"{customer} 맞춤 제안서"
        self.over = {}; self.n = 0
        self.sections = []      # [(장 제목, 시작 쪽)] — deck.slide(section=...) 로 기록, 목차에 쓰임
        self._toc = None
        self.issues = []        # 형식 규칙 위반 (표지 제목, 목차 누락 등)

    def _frame(self, s, cover=False):
        """모든 슬라이드 공통: 상단 우측 고객사명, 하단 우측 로고, 하단 좌측 쪽번호."""
        sl = Slide(self, s, self.n)
        sl._tb(W - MX - 3.6, 0.32, 3.6, 0.28, self.customer, 10.5, NAVY, True, align=PP_ALIGN.RIGHT, name="header_customer")
        lh = 0.27
        from PIL import Image
        with Image.open(self.logo) as im: ratio = im.width / im.height
        s.shapes.add_picture(self.logo, Inches(W - MX - lh * ratio), Inches(H - 0.33 - lh), Inches(lh * ratio), Inches(lh)).name = "footer_logo"
        if not cover:
            sl._tb(MX, H - 0.6, 4.2, 0.26, f"{self.n:02d}   {self.footer}", 8.5, GRAY, name="footer_page")
        return sl

    def cover(self, title, subtitle, date, proposer, contact=None):
        """표지. title 은 반드시 'ㅇㅇㅇ 제안서' 형식 (예: '외부 협업 파일 보안 공유 체계 구축 제안서').
        표지에 '○○ 귀중'·'수신:' 줄은 넣지 않는다 — 고객사명은 모든 슬라이드 상단 우측에 이미 있다."""
        if not title.replace(" ", "").endswith("제안서"):
            self.issues.append(f"표지 제목이 'ㅇㅇㅇ 제안서' 형식이 아님: {title!r}")
        self.n += 1
        s = self.p.slides.add_slide(self.p.slide_layouts[6])
        sl = self._frame(s, cover=True)
        # 표지 모티프: 오른쪽 위에서 들어오는 큰 앰버 원 (부분만 보이게)
        o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(W - 2.6), Inches(-1.6), Inches(4.2), Inches(4.2))
        o.fill.solid(); o.fill.fore_color.rgb = _rgb(TINT); o.line.fill.background(); o.shadow.inherit = False; o.name = "cover_motif"
        s.shapes._spTree.remove(o._element); s.shapes._spTree.insert(2, o._element)   # 맨 뒤로
        th = est_h(title, CW, 30, 1.2, bold=True)
        sl._tb(MX, 3.3, CW, th, title, 30, NAVY, True, spacing=1.2, name="cover_title")
        sh = est_h(subtitle, CW, 14, 1.35)
        sl._tb(MX, 3.5 + th, CW, sh, subtitle, 14, INK, spacing=1.35, name="cover_subtitle")
        y = H - 2.45
        sl._tb(MX, y, CW, 0.3, date, 12, INK)
        sl._tb(MX, y + 0.36, CW, 0.32, proposer, 13.5, NAVY, True)
        if contact: sl._tb(MX, y + 0.74, CW, est_h(contact, CW, 10), contact, 10, GRAY)
        return sl

    def toc(self, title="제안 목차", kicker="CONTENTS"):
        """표지 다음 장에 둔다. 목차 행(장 번호·장 제목·쪽)은 save() 때 section 기록으로 채워진다."""
        if self.n != 1:
            self.issues.append("제안 목차는 표지 바로 다음(2쪽)에 와야 함")
        self._toc = self.slide(title, kicker=kicker)
        return self._toc

    def _draw_toc(self):
        sl = self._toc; n = len(self.sections)
        if not n:
            self.issues.append("목차에 넣을 section 이 없음 — deck.slide(..., section='제안 배경 및 필요성') 처럼 지정"); return
        avail = BOTTOM - sl.y; gap = 0.12
        h = min(0.62, (avail - gap * (n - 1)) / n)
        num = 0
        for i, (name, page) in enumerate(self.sections):
            y = sl.y + i * (h + gap)
            appendix = name.startswith("별첨") or name.lower().startswith("appendix")
            if not appendix: num += 1
            sl._box(MX, y, CW, h, fill=SOFT, radius=0.18, name=f"toc{i+1}")
            sl._tb(MX + 0.2, y, 0.7, h, "별첨" if appendix else f"{num:02d}", 15 if not appendix else 12,
                   AMBER_DK, True, anchor=MSO_ANCHOR.MIDDLE)
            sl._tb(MX + 0.95, y, CW - 2.0, h, name.replace("별첨 — ", "").replace("별첨 ", "") if appendix else name,
                   13.5, NAVY, True, anchor=MSO_ANCHOR.MIDDLE)
            sl._tb(MX + CW - 1.0, y, 0.8, h, f"{page:02d}", 13, GRAY, True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

    def slide(self, title, kicker=None, lead=None, section=None):
        """본문 슬라이드: kicker(작은 앰버 라벨) + 제목 + (선택) 리드 문장.
        section: 이 슬라이드가 속한 장 제목(목차 항목). 장이 바뀌는 첫 슬라이드의 쪽번호가 목차에 들어간다."""
        self.n += 1
        if section and (not self.sections or self.sections[-1][0] != section):
            self.sections.append((section, self.n))
        s = self.p.slides.add_slide(self.p.slide_layouts[6])
        sl = self._frame(s)
        y = 0.72
        if kicker:
            sl._tb(MX, y, CW, 0.24, kicker, 9.5, AMBER_DK, True); y += 0.28
        th = est_h(title, CW, 22, 1.18, bold=True)
        sl._tb(MX, y, CW, th, title, 22, NAVY, True, spacing=1.18, name="title"); y += th + 0.14
        if lead:
            lh = est_h(lead, CW, 11.5, 1.4)
            sl._tb(MX, y, CW, lh, lead, 11.5, GRAY, spacing=1.4); y += lh + 0.12
        sl.y = y + 0.12
        return sl

    def save(self, path):
        if self._toc is not None:
            self._draw_toc()
        else:
            self.issues.append("제안 목차 슬라이드 없음 — 표지 다음에 deck.toc()")
        self.p.save(path)
        try: sys.stdout.reconfigure(encoding="utf-8")
        except Exception: pass
        for msg in self.issues:
            print("WARNING", msg)
        for no, over in sorted(self.over.items()):
            print(f"WARNING slide {no}: 하단 여백을 {over:.2f}in 넘침 - 줄이거나 슬라이드를 나눌 것")
        print(f"saved {path} ({self.n} slides, {len(self.over) + len(self.issues)} warnings)")
        return path

    @property
    def warnings(self):
        return [f"slide {k}" for k in sorted(self.over)]
