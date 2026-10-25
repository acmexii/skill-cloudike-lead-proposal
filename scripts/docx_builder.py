# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# 출처: acmexii/skill-rnd-proposal assets/docx_builder.py — 영업 제안서용으로 폰트·표지·콜아웃만 추가
BODY_FONT="맑은 고딕"; BODY_SZ=10.5; TBL_SZ=8.8
DIA=os.path.join(os.path.dirname(os.path.abspath(__file__)),"dia")

def _rpr_font(run, name=BODY_FONT, size=BODY_SZ, bold=False, color=None):
    run.font.name=name; run.font.size=Pt(size); run.bold=bold
    if color: run.font.color.rgb=RGBColor.from_string(color)
    rPr=run._element.get_or_add_rPr(); rf=rPr.find(qn('w:rFonts'))
    if rf is None:
        rf=OxmlElement('w:rFonts'); rPr.append(rf)
    for a in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'): rf.set(qn(a), name)

def shade(cell, hexcolor):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd')
    sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),hexcolor)
    tcPr.append(sh)

def cell_border(cell, sz=4, color="9AA0A6"):
    tcPr=cell._tc.get_or_add_tcPr(); b=OxmlElement('w:tcBorders')
    for e in ('top','left','bottom','right'):
        el=OxmlElement('w:'+e); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),str(sz))
        el.set(qn('w:space'),'0'); el.set(qn('w:color'),color); b.append(el)
    tcPr.append(b)

def para_border(p, color="333A41", sz=10, bottom_only=False):
    pPr=p._p.get_or_add_pPr(); b=OxmlElement('w:pBdr')
    edges=['bottom'] if bottom_only else ['top','left','bottom','right']
    for e in edges:
        el=OxmlElement('w:'+e); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),str(sz))
        el.set(qn('w:space'),'3'); el.set(qn('w:color'),color); b.append(el)
    pPr.append(b)

def para_shade(p, hexcolor):
    pPr=p._p.get_or_add_pPr(); sh=OxmlElement('w:shd')
    sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),hexcolor); pPr.append(sh)

def rich(p, text, size=BODY_SZ, base_bold=False, color=None):
    """**bold** 지원"""
    parts=text.split("**")
    for i,seg in enumerate(parts):
        if not seg: continue
        r=p.add_run(seg); _rpr_font(r, BODY_FONT, size, base_bold or (i%2==1), color)
    return p

class Doc:
    def __init__(self):
        self.d=Document()
        s=self.d.sections[0]
        s.page_width=Cm(21.0); s.page_height=Cm(29.7)
        s.left_margin=Cm(2.0); s.right_margin=Cm(2.0); s.top_margin=Cm(2.0); s.bottom_margin=Cm(1.8)
        st=self.d.styles['Normal']; st.font.name=BODY_FONT; st.font.size=Pt(BODY_SZ)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), BODY_FONT)
        st.paragraph_format.space_after=Pt(2); st.paragraph_format.line_spacing=1.24
        self.CW=Cm(17.0)
    # ---------- headings
    def h1(self, text):
        p=self.d.add_paragraph(); p.paragraph_format.space_before=Pt(16); p.paragraph_format.space_after=Pt(8)
        p.paragraph_format.keep_with_next=True
        rich(p, text, 14.5, True); para_border(p,"1F3864",14); para_shade(p,"E8EDF5")
        return p
    def h2(self, text):
        p=self.d.add_paragraph(); p.paragraph_format.space_before=Pt(13); p.paragraph_format.space_after=Pt(5)
        p.paragraph_format.keep_with_next=True
        rich(p, "▣ "+text, 12.5, True, "1F3864"); para_border(p,"1F3864",6,bottom_only=True)
        return p
    def h3(self, text):
        p=self.d.add_paragraph(); p.paragraph_format.space_before=Pt(9); p.paragraph_format.space_after=Pt(3)
        p.paragraph_format.keep_with_next=True
        rich(p, text, 11.5, True, "2E5E8C"); return p
    def h4(self, text):
        p=self.d.add_paragraph(); p.paragraph_format.space_before=Pt(7); p.paragraph_format.space_after=Pt(2)
        p.paragraph_format.keep_with_next=True
        rich(p, "□ "+text, 11.0, True); return p
    # ---------- body
    def o(self, text, size=BODY_SZ):
        p=self.d.add_paragraph(); p.paragraph_format.left_indent=Cm(0.45); p.paragraph_format.first_line_indent=Cm(-0.45)
        p.paragraph_format.space_before=Pt(3)
        rich(p, "o "+text, size); return p
    def b(self, text, size=BODY_SZ-0.5):
        p=self.d.add_paragraph(); p.paragraph_format.left_indent=Cm(0.9); p.paragraph_format.first_line_indent=Cm(-0.4)
        rich(p, "- "+text, size); return p
    def b2(self, text, size=BODY_SZ-0.5):
        p=self.d.add_paragraph(); p.paragraph_format.left_indent=Cm(1.35); p.paragraph_format.first_line_indent=Cm(-0.35)
        rich(p, "· "+text, size); return p
    def plain(self, text, size=BODY_SZ, indent=0.0, align=None, before=0, after=0):
        p=self.d.add_paragraph(); p.paragraph_format.left_indent=Cm(indent)
        p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after)
        if align: p.alignment=align
        rich(p, text, size); return p
    def note(self, text, size=8.5):
        p=self.d.add_paragraph(); p.paragraph_format.left_indent=Cm(0.5)
        rich(p, "※ "+text, size, color="4A5158"); return p
    def caption(self, text, size=9.0):
        p=self.d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(8)
        rich(p, text, size, color="3A4148"); return p
    def tcap(self, text, size=10.0):
        p=self.d.add_paragraph(); p.paragraph_format.space_before=Pt(9); p.paragraph_format.space_after=Pt(3)
        p.paragraph_format.keep_with_next=True
        rich(p, "< "+text+" >", size, True, "1F3864"); return p
    def fig(self, name, caption=None, width_cm=17.0):
        p=self.d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(2)
        p.add_run().add_picture(os.path.join(DIA,name+".png"), width=Cm(width_cm))
        if caption: self.caption(caption)
    def cover(self, customer, title, subtitle, date, proposer, contact=None, recipient=None):
        """표지 1쪽: 고객사 귀중(+수신: 담당자) / 제목 / 부제 / 일자·제안사. 끝에 쪽나눔."""
        self.plain("", before=90)
        self.plain(f"{customer} 귀중", 13, align=WD_ALIGN_PARAGRAPH.LEFT, after=2 if recipient else 10)
        if recipient: self.plain(f"수신: {recipient}", 10.5, after=10)
        p=self.d.add_paragraph(); rich(p, title, 24, True, "1F3864"); para_border(p,"1F3864",18,bottom_only=True)
        p.paragraph_format.space_after=Pt(10)
        self.plain(subtitle, 12.5, after=120)
        self.plain(date, 11, after=2); self.plain(f"**{proposer}**", 12, after=2)
        if contact: self.plain(contact, 9.5)
        self.pagebreak()
    def callout(self, text, size=10.5, fill="EEF3FA", color="1F3864"):
        """한 단락 강조 박스 (핵심 제안 명제, 한 장 요약용)."""
        p=self.d.add_paragraph(); p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(8)
        p.paragraph_format.left_indent=Cm(0.2); p.paragraph_format.right_indent=Cm(0.2)
        rich(p, text, size, False, color); para_border(p, color, 8); para_shade(p, fill); return p
    def pagebreak(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    def table(self, rows, widths=None, header=True, size=TBL_SZ, align_center_cols=(), hdr_fill="DCE3EE", zebra=None):
        n=max(len(r) for r in rows)
        rows=[list(r)+['']*(n-len(r)) for r in rows]
        t=self.d.add_table(rows=len(rows), cols=n)
        t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
        tblPr=t._tbl.tblPr; lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tblPr.append(lay)
        if widths:
            widths=list(widths)+[widths[-1]]*(n-len(widths)) if len(widths)<n else widths[:n]
            tot=sum(widths); ws=[Cm(17.0*w/tot) for w in widths]
        else:
            ws=[Cm(17.0/n)]*n
        # explicit table width + grid (LibreOffice/Word honour the grid, not per-cell tcW alone)
        tw=OxmlElement('w:tblW'); tw.set(qn('w:w'), str(int(Cm(17.0).twips))); tw.set(qn('w:type'),'dxa'); tblPr.append(tw)
        cm=OxmlElement('w:tblCellMar')
        for side,val in (('top','20'),('bottom','20'),('left','60'),('right','60')):
            e=OxmlElement('w:'+side); e.set(qn('w:w'),val); e.set(qn('w:type'),'dxa'); cm.append(e)
        tblPr.append(cm)
        grid=t._tbl.find(qn('w:tblGrid'))
        if grid is not None: t._tbl.remove(grid)
        grid=OxmlElement('w:tblGrid')
        for wv in ws:
            gc=OxmlElement('w:gridCol'); gc.set(qn('w:w'), str(int(wv.twips))); grid.append(gc)
        t._tbl.insert(list(t._tbl).index(tblPr)+1, grid)
        for ri,row in enumerate(rows):
            for ci,val in enumerate(row):
                c=t.cell(ri,ci); c.width=ws[ci]
                p=c.paragraphs[0]; p.paragraph_format.space_before=Pt(1); p.paragraph_format.space_after=Pt(1)
                p.paragraph_format.line_spacing=1.02
                txt="" if str(val) in ("<<","^^") else str(val)
                bold = (header and ri==0)
                if ci in align_center_cols or (header and ri==0): p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                lines=txt.split("\n")
                for li,ln in enumerate(lines):
                    if li>0:
                        p=c.add_paragraph(); p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0)
                        p.paragraph_format.line_spacing=1.02
                        if ci in align_center_cols or (header and ri==0): p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    rich(p, ln, size, bold)
                cell_border(c)
                if header and ri==0:
                    shade(c, hdr_fill)
                    for hp in c.paragraphs: hp.paragraph_format.keep_with_next=True  # 머리행이 쪽 끝에 홀로 남지 않게
                elif zebra and ri%2==0: shade(c, zebra)
            # 행이 쪽 경계에서 쪼개지지 않게, 머리행은 다음 쪽에서 반복
            trPr=t.rows[ri]._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
            if header and ri==0: trPr.append(OxmlElement('w:tblHeader'))
        # horizontal merge: '<<' merges into the cell on the left; '^^' merges into the cell above
        for ri,row in enumerate(rows):
            for ci,val in enumerate(row):
                if val=='<<' and ci>0:
                    try: t.cell(ri,ci-1).merge(t.cell(ri,ci))
                    except Exception: pass
        for ri,row in enumerate(rows):
            for ci,val in enumerate(row):
                if val=='^^' and ri>0:
                    try: t.cell(ri-1,ci).merge(t.cell(ri,ci))
                    except Exception: pass
        self.d.add_paragraph().paragraph_format.space_after=Pt(0)
        return t
    def footer(self, text):
        """모든 쪽 하단 가운데에 '<text> · 쪽번호'. 표지(첫 쪽)에는 표시하지 않음."""
        sec=self.d.sections[0]; sec.different_first_page_header_footer=True
        p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        rich(p, text+" · ", 8, color="4A5158")
        for kind,val in (('begin',None),(None,'PAGE'),('end',None)):
            r=p.add_run(); _rpr_font(r, BODY_FONT, 8, color="4A5158")
            if kind:
                e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), kind)
            else:
                e=OxmlElement('w:instrText'); e.set(qn('xml:space'),'preserve'); e.text=val
            r._r.append(e)
    def save(self, path):
        self.d.save(path); return path
