# -*- coding: utf-8 -*-
"""리드 한 건의 참조 자료(녹취록·회의록·메일·메모·견적·RFP 등)를 찾아 목록과 텍스트를 만든다.

리드 관리 대장만 있으면 제안서는 작성된다. 이 스크립트는 '있으면 쓰는' 자료를 빠짐없이 모으기 위한 것이다.

찾는 곳 (대장 파일이 있는 폴더 = BASE 기준)
  1) 리드 폴더: BASE 아래(깊이 2까지)에서 폴더 이름에 고객사 핵심 이름이 들어간 폴더 전체
     예) BASE/에이아이랩스/,  BASE/리드자료/(주)에이아이랩스/,  BASE/20261020_에이아이랩스/
  2) BASE 바로 아래 파일 중 파일명에 고객사 핵심 이름이 들어간 것
     예) BASE/녹취록_에이아이랩스_도입미팅_20261025.txt
  --folder 로 리드 폴더를 직접 지정할 수도 있다.

예)  python find_lead_materials.py "<대장.xlsx>" --company "(주)에이아이랩스"
     python find_lead_materials.py "<대장.xlsx>" --company 대한바이오 --extract <작업폴더>/materials
     (--extract: docx/pdf/pptx/xlsx/txt 의 텍스트를 .txt 로 뽑아 저장 → Read 로 읽는다)
"""
import argparse, json, os, re, sys

# 파일명 키워드 → 자료 유형. 위에서부터 먼저 맞는 것.
KINDS = [
    ("녹취록", ["녹취", "transcript", "stt", "통화"]),
    ("회의록", ["회의록", "미팅록", "미팅노트", "meeting", "minutes", "회의"]),
    ("요구사항·RFP", ["rfp", "제안요청", "요구사항", "요청서", "체크리스트", "보안점검", "가이드라인"]),
    ("메일", ["메일", "mail", "email", ".eml", ".msg"]),
    ("기존 제안서·견적", ["제안서", "견적", "quote", "proposal"]),
    ("메모", ["메모", "memo", "note", "노트"]),
    ("고객 자료", ["회사소개", "조직도", "현황", "소개서"]),
]
TEXT_EXT = {".txt", ".md", ".vtt", ".srt", ".csv", ".eml"}
DOC_EXT = {".docx", ".pdf", ".pptx", ".xlsx", ".hwpx"}
MEDIA_EXT = {".mp4", ".m4a", ".mp3", ".wav", ".mov", ".webm"}
IMG_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
SKIP_PREFIX = ("~$", ".")
# 이 스킬이 만든 산출물 폴더(페이지 PNG, 작업 파일)는 참조 자료가 아니다
SKIP_DIR = ("_pages", "_work", "materials", "__pycache__")


def skip_dir(name):
    return name.endswith(SKIP_DIR) or name.startswith(".")


def core_name(company):
    """'(주)에이아이랩스', '주식회사 에이아이랩스' → '에이아이랩스'"""
    s = re.sub(r"\(.*?\)|㈜|주식회사|유한회사|\s", "", company)
    return s or company


def kind_of(name):
    low = name.lower()
    for k, kws in KINDS:
        if any(kw in low for kw in kws):
            return k
    ext = os.path.splitext(low)[1]
    if ext in MEDIA_EXT:
        return "음성·영상"
    return "기타"


def find(base, core, folder=None):
    hits, dirs = [], []
    if folder:
        dirs.append(os.path.abspath(folder))
    else:
        for root, ds, _ in os.walk(base):
            depth = os.path.relpath(root, base).count(os.sep) + (0 if root == base else 1)
            if depth >= 2:
                ds[:] = []
            ds[:] = [d for d in ds if not skip_dir(d)]
            for d in ds:
                if core in d.replace(" ", ""):
                    dirs.append(os.path.join(root, d))
    seen = set()
    for d in dirs:
        for root, sub, fs in os.walk(d):
            sub[:] = [x for x in sub if not skip_dir(x)]
            for f in fs:
                if not f.startswith(SKIP_PREFIX):
                    seen.add(os.path.join(root, f))
    if not folder:
        for f in os.listdir(base):
            p = os.path.join(base, f)
            if os.path.isfile(p) and core in f.replace(" ", "") and not f.startswith(SKIP_PREFIX):
                seen.add(p)
    for p in sorted(seen):
        st = os.stat(p)
        ext = os.path.splitext(p)[1].lower()
        hits.append(dict(path=p, kind=kind_of(os.path.basename(p)), ext=ext, kb=round(st.st_size / 1024),
                         modified=__import__("datetime").datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d"),
                         readable="text" if ext in TEXT_EXT else "extract" if ext in DOC_EXT
                         else "image" if ext in IMG_EXT
                         else "transcribe-needed" if ext in MEDIA_EXT else "skip"))
    # 같은 이름의 텍스트 전사본이 있는 음성·영상은 전사본으로 대신한다
    stems = {os.path.splitext(h["path"])[0] for h in hits if h["readable"] == "text"}
    for h in hits:
        if h["readable"] == "transcribe-needed" and os.path.splitext(h["path"])[0] in stems:
            h["readable"] = "covered-by-transcript"
    return dirs, hits


def extract_text(p):
    ext = os.path.splitext(p)[1].lower()
    if ext in TEXT_EXT:
        for enc in ("utf-8-sig", "cp949"):
            try:
                return open(p, encoding=enc).read()
            except UnicodeDecodeError:
                pass
    if ext == ".docx":
        import docx
        d = docx.Document(p)
        out = [x.text for x in d.paragraphs if x.text.strip()]
        for t in d.tables:
            out += [" | ".join(c.text.strip() for c in r.cells) for r in t.rows]
        return "\n".join(out)
    if ext == ".pdf":
        import pypdf
        return "\n".join((pg.extract_text() or "") for pg in pypdf.PdfReader(p).pages)
    if ext == ".pptx":
        from pptx import Presentation
        out = []
        for i, s in enumerate(Presentation(p).slides, 1):
            out.append(f"--- slide {i}")
            out += [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        return "\n".join(out)
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
        out = []
        for ws in wb:
            out.append(f"--- sheet {ws.title}")
            out += [" | ".join("" if c is None else str(c) for c in r) for r in ws.iter_rows(values_only=True)
                    if any(c is not None for c in r)]
        return "\n".join(out)
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("ledger", help="리드 관리 대장 경로 (그 폴더가 검색 기준)")
    ap.add_argument("--company", required=True)
    ap.add_argument("--folder", help="리드 폴더를 직접 지정")
    ap.add_argument("--extract", help="텍스트 추출본을 저장할 폴더")
    a = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")
    base = os.path.dirname(os.path.abspath(a.ledger))
    core = core_name(a.company)
    dirs, hits = find(base, core, a.folder)
    hits = [h for h in hits if os.path.abspath(h["path"]) != os.path.abspath(a.ledger)]
    if a.extract:
        os.makedirs(a.extract, exist_ok=True)
        for h in hits:
            if h["readable"] in ("text", "extract"):
                try:
                    txt = extract_text(h["path"])
                except Exception as e:  # 손상·암호화 파일은 목록에만 남긴다
                    h["extract_error"] = str(e)[:120]
                    continue
                if txt:
                    out = os.path.join(a.extract, os.path.basename(h["path"]) + ".txt")
                    open(out, "w", encoding="utf-8").write(txt)
                    h["text"] = out
    print(json.dumps(dict(company=a.company, search_key=core, base=base, lead_folders=dirs,
                          materials=hits, count=len(hits),
                          mode="ledger+materials" if hits else "ledger-only"),
                     ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
